"""Чтение счёта на оплату из Excel (выгрузки 1С и похожих программ).

Печатные формы счёта бывают разными, поэтому ничего не ищется по фиксированным адресам:
- шапка таблицы товаров находится по синонимам заголовков («Товары», «Наименование»,
  «Кол-во», «Количество», «Ед.», «Цена», «Сумма»…);
- позиции — строки под шапкой, где есть наименование и число;
- покупатель разбирается при любом порядке частей (название, БИН, адрес);
- итоги и НДС — по смыслу подписи («Итого», «Всего к оплате», «В т.ч. НДС», «В том числе НДС»…).
"""
from __future__ import annotations

import re
import tempfile
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import openpyxl

MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
    "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
}

# Роли колонок таблицы и их возможные заголовки (сравнение — по началу текста, без регистра)
HEADER_SYNONYMS = {
    "no": ("№", "n", "#", "№ п/п", "п/п", "nn"),
    "code": ("код", "артикул", "арт."),
    "name": ("товар", "наименование", "номенклатура", "работ", "услуг", "описание"),
    "qty": ("кол-во", "кол.", "количество", "кол "),
    "unit": ("ед", "ед.", "ед. изм", "единица"),
    "price": ("цена",),
    "sum": ("сумма", "стоимость", "всего"),
}

LEGAL_FORMS = ("товарищество", "тоо", "ао ", "ао«", "ао\"", "акционерное", "ип ", "индивидуальный предприниматель",
               "ооо", "общество", "тдо", "кх ", "крестьянское", "гу ", "кгп", "ргп", "учреждение", "фонд",
               "филиал", "представительство")


class InvoiceError(Exception):
    pass


@dataclass
class Item:
    name: str
    qty: float
    unit: str
    price: float
    total: float


@dataclass
class Invoice:
    number: str = ""
    date: date | None = None
    buyer: str = ""
    buyer_details: str = ""
    items: list[Item] = field(default_factory=list)
    total: float = 0.0
    vat: float | None = None
    vat_label: str = "В т.ч. НДС:"
    source: str = ""


# ---------------------------------------------------------------- мелкие помощники
def _num(v) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.replace("\xa0", "").replace(" ", "").replace(",", ".").strip()
        if re.fullmatch(r"-?\d+(\.\d+)?", s):
            return float(s)
    return None


def _norm(v) -> str:
    return " ".join(str(v).replace("\xa0", " ").split()).strip() if v is not None else ""


def _parse_date(text: str) -> date | None:
    m = re.search(r"(\d{1,2})\s+([а-яё]+)\s+(\d{4})", text, re.I)
    if m and m.group(2).lower() in MONTHS:
        return date(int(m.group(3)), MONTHS[m.group(2).lower()], int(m.group(1)))
    m = re.search(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})", text)
    if m:
        y = int(m.group(3))
        return date(y + 2000 if y < 100 else y, int(m.group(2)), int(m.group(1)))
    return None


def split_buyer(text: str) -> tuple[str, str]:
    """Разобрать строку покупателя в (название, реквизиты) при любом порядке частей:
    «ТОО "Ромашка", БИН / ИИН 123, г.Алматы» и «БИН / ИИН 123,Товарищество … "Ромашка",г.Алматы»."""
    text = _norm(text).strip(" ,;")
    bin_part = ""
    m = re.search(r"(БИН\s*/\s*ИИН|ИИН\s*/\s*БИН|БИН|ИИН)\s*:?\s*(\d{12})", text, re.I)
    if m:
        bin_part = f"{m.group(1).upper().replace(' ', '').replace('/', ' / ')} {m.group(2)}"
        text = (text[:m.start()] + text[m.end():]).strip(" ,;")
    parts = [p.strip() for p in re.split(r",(?=(?:[^\"«»]*[\"«][^\"«»]*[\"»])*[^\"«»]*$)", text) if p.strip()]
    name_idx = None
    for i, p in enumerate(parts):
        low = p.lower() + " "
        if any(low.startswith(f) or f" {f}" in f" {low}" for f in LEGAL_FORMS) or re.search(r"[\"«].+[\"»]", p):
            name_idx = i
            break
    if name_idx is None:
        name_idx = 0 if parts else None
    name = parts[name_idx] if name_idx is not None else ""
    rest = [p for i, p in enumerate(parts) if i != name_idx]
    details = ", ".join(([bin_part] if bin_part else []) + rest)
    return name, details


def _convert_xls(path: Path) -> Path:
    """Старый .xls через установленный Excel пересохраняем в .xlsx."""
    try:
        import win32com.client  # noqa: WPS433
    except ImportError as e:
        raise InvoiceError("Для файлов .xls нужен установленный Microsoft Excel.") from e
    out = Path(tempfile.mkdtemp()) / (path.stem + ".xlsx")
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(path.resolve()), ReadOnly=True)
        wb.SaveAs(str(out), FileFormat=51)
        wb.Close(False)
    finally:
        excel.Quit()
    return out


# ---------------------------------------------------------------- разбор листа
class _Grid:
    def __init__(self, ws):
        self.cells: dict[tuple[int, int], object] = {}
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and _norm(c.value) != "":
                    self.cells[(c.row, c.column)] = c.value
        self.max_row = max((r for r, _ in self.cells), default=0)

    def row(self, r: int) -> list[tuple[int, object]]:
        return sorted((c, v) for (rr, c), v in self.cells.items() if rr == r)

    def text_cells(self):
        return [(r, c, _norm(v)) for (r, c), v in sorted(self.cells.items()) if isinstance(v, str)]


def _role(text: str) -> str | None:
    t = text.lower().strip().rstrip(":")
    if not t:
        return None
    if t in ("№", "n", "#", "nn", "№п/п", "№ п/п", "п/п"):
        return "no"
    for role in ("qty", "unit", "price", "sum", "code", "name"):
        for syn in HEADER_SYNONYMS[role]:
            if t == syn.strip() or t.startswith(syn):
                if role == "unit" and t.startswith("единиц") is False and len(t) > 12:
                    continue
                return role
    return None


def _find_header(g: _Grid) -> tuple[int, dict[str, int]] | None:
    """Строка шапки с наибольшим числом распознанных колонок (нужны наименование и кол-во/сумма).
    Учитываем шапку в две строки: роли из соседней строки добавляются, если колонка не занята."""
    best = None
    for r in range(1, g.max_row + 1):
        cols: dict[str, int] = {}
        for c, v in g.row(r):
            if isinstance(v, str):
                role = _role(_norm(v))
                if role and role not in cols:
                    cols[role] = c
        if "name" not in cols or not ({"qty", "sum"} & cols.keys()):
            continue
        for c, v in g.row(r + 1):  # вторая строка шапки (например «Кол-во | Ед.» под «Количество»)
            if isinstance(v, str):
                role = _role(_norm(v))
                if role and role not in cols and c not in cols.values():
                    cols[role] = c
        score = len(cols)
        if not best or score > best[2]:
            best = (r, cols, score)
    return (best[0], best[1]) if best else None


def _parse_sheet(ws, source: str) -> Invoice | None:
    g = _Grid(ws)
    found = _find_header(g)
    if not found:
        return None
    header_row, cols = found
    inv = Invoice(source=source)

    # номер и дата
    for r, c, t in g.text_cells():
        m = re.search(r"Сч[её]т(?:\s+на\s+оплату)?\s*№\s*([^\s,]+)\s*(?:от\s+(.+))?", t, re.I)
        if m:
            inv.number = m.group(1)
            inv.date = _parse_date(m.group(2) or "")
            break

    # покупатель — первое значение справа от «Покупатель», «Заказчик», «Плательщик», «Кому»
    for r, c, t in g.text_cells():
        if re.match(r"^(покупатель|заказчик|плательщик|получатель услуг|кому)\b", t, re.I):
            after = re.sub(r"^[^:]*:\s*", "", t) if ":" in t else ""
            rest = [_norm(v) for col, v in g.row(r) if col > c and isinstance(v, str) and _norm(v)]
            value = after or (rest[0] if rest else "")
            if value:
                inv.buyer, inv.buyer_details = split_buyer(value)
                break

    # колонки, которые считаются числовыми, — чтобы единица измерения не попала из них
    numeric_cols = {cols[k] for k in ("no", "qty", "price", "sum", "code") if k in cols}
    name_col = cols["name"]
    started, blank_run, r = False, 0, header_row + 1
    while r <= g.max_row:
        row = dict(g.row(r))
        texts = " ".join(_norm(v) for v in row.values() if isinstance(v, str)).lower()
        if started and re.search(r"\b(итого|всего|в т\.ч|в том числе)\b", texts):
            break
        name = _norm(row.get(name_col, ""))
        qty = _num(row.get(cols.get("qty")))
        total = _num(row.get(cols.get("sum")))
        price = _num(row.get(cols.get("price")))
        if name and (qty is not None or total is not None) and _num(name) is None:
            started, blank_run = True, 0
            unit = _norm(row.get(cols["unit"], "")) if "unit" in cols else ""
            if not unit and "qty" in cols:
                # единица без своего заголовка — текст справа от количества, до цены
                limit = cols.get("price", cols.get("sum", 10 ** 6))
                unit = next((_norm(v) for col, v in sorted(row.items())
                             if cols["qty"] < col < limit and isinstance(v, str) and col not in numeric_cols), "")
            qty = qty if qty is not None else 1.0
            if price is None:
                price = round(total / qty, 2) if total is not None and qty else 0.0
            inv.items.append(Item(name=name, qty=qty, unit=unit, price=price,
                                  total=total if total is not None else round(qty * price, 2)))
        elif started:
            blank_run += 1
            if blank_run >= 2:
                break
        r += 1

    if not inv.items:
        return None

    # итоги
    totals: dict[str, float] = {}
    for rr in range(r, g.max_row + 1):
        row = g.row(rr)
        labels = [(c, _norm(v)) for c, v in row if isinstance(v, str) and _num(v) is None]
        for c, label in labels:
            nums = [_num(v) for col, v in row if col > c and _num(v) is not None]
            if not nums:
                continue
            low = label.lower()
            if "ндс" in low and "с ндс" not in low and "без ндс" not in low:
                inv.vat = nums[-1]
                inv.vat_label = re.sub(r"(?i)в том числе", "В т.ч.", label)
            elif re.search(r"итого с ндс|всего к оплате|итого к оплате", low):
                totals["with_vat"] = nums[-1]
            elif low.startswith("итого") or low.startswith("всего"):
                totals.setdefault("total", nums[-1])
            break
    inv.total = totals.get("with_vat") or totals.get("total") or round(sum(i.total for i in inv.items), 2)
    return inv


def parse_invoice(path: str | Path) -> Invoice:
    path = Path(path)
    if path.suffix.lower() == ".xls":
        path = _convert_xls(path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise InvoiceError(f"Не удалось открыть файл:\n{e}") from e
    for ws in wb.worksheets:
        inv = _parse_sheet(ws, str(path))
        if inv:
            return inv
    raise InvoiceError("В файле не найдена таблица товаров.\n"
                       "Нужна таблица с колонками наименования и количества или суммы "
                       "(например: №, Наименование / Товары, Кол-во, Ед., Цена, Сумма).")
