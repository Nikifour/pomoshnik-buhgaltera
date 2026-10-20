"""Сборка КП в Excel по макету из блоков (core/layout.py).

Каждый тип блока рисует себя функцией _b_<тип>; порядок и настройки блоков задаёт
пользователь в конструкторе, а вёрстка внутри блоков — здесь.
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import column_index_from_string
from openpyxl.utils.units import pixels_to_EMU
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

from core.invoice_parser import Invoice, Item
from core import richtext
from core.layout import normalize
from core.num_words import amount_in_words

FONT = "Calibri"
GREY = "6B7280"
DARK = "1F2937"
LINE = "E5E7EB"

# Колонки: поле, №, наименование, кол-во, ед., цена, сумма, поле
COLS = {"A": 1.5, "B": 5, "C": 50, "D": 9, "E": 7, "F": 14, "G": 16, "H": 1.5}
TABLE_WIDTH = sum(COLS[c] for c in "BCDEFG")
FIRST, LAST = "B", "G"
NL_CHAR = chr(10)


@dataclass
class KPMeta:
    number: str  # по умолчанию — номер счёта на оплату
    date: date
    buyer: str
    buyer_details: str = ""


def _hex(color: str) -> str:
    return color.lstrip("#").upper()


def _tint(color: str, factor: float) -> str:
    """Осветлить цвет (factor 0..1, ближе к 1 — светлее)."""
    c = _hex(color)
    r, g, b = (int(c[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda x: round(x + (255 - x) * factor)  # noqa: E731
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


def _lines(text: str, width_chars: float) -> int:
    return sum(max(1, math.ceil(len(p) / max(width_chars, 1))) for p in text.split("\n"))


# Метка значка WhatsApp в тексте: {WhatsApp}, {WA}, {Вотсап}, {Ватсап}
WA_TOKEN = re.compile(r"\{(?:whatsapp|wa|вотсап|ватсап)\}", re.I)
_TENGE = re.compile(r"\b(т)(енге)\b", re.I)
_FONT_FILES = {(False, False): "calibri.ttf", (True, False): "calibrib.ttf",
               (False, True): "calibrii.ttf", (True, True): "calibriz.ttf"}
_font_cache: dict = {}


def legal_tenge(text: str) -> str:
    """«тенге» -> «теңге» (так требует закон), с сохранением регистра."""
    def repl(m):
        t, rest = m.group(1), m.group(2)
        ng = "Ң" if rest[0].isupper() else "ң"
        return f"{t}{rest[0]}{ng}{rest[2:]}"
    return _TENGE.sub(repl, text)


def _text_px(text: str, size_pt: float, bold=False, italic=False) -> float:
    """Ширина текста в пикселях так, как её нарисует Excel (Calibri, 96 dpi)."""
    from PIL import ImageFont
    px = size_pt * 96 / 72
    key = (bold, italic, round(px * 4))
    if key not in _font_cache:
        try:
            _font_cache[key] = ImageFont.truetype(f"C:/Windows/Fonts/{_FONT_FILES[(bold, italic)]}", px * 4)
        except OSError:
            _font_cache[key] = None
    f = _font_cache[key]
    return f.getlength(text) / 4 if f else len(text) * px * 0.5


def _fmt_money(v: float) -> str:
    return f"{v:,.2f}".replace(",", " ").replace(".", ",")


def _fmt_qty(v: float) -> float | int:
    return int(v) if float(v).is_integer() else v


def fill_placeholders(text: str, s: dict, meta: KPMeta) -> str:
    return (text.replace("{КОМПАНИЯ}", s["company_name"])
                .replace("{ПОКУПАТЕЛЬ}", meta.buyer)
                .replace("{МЕНЕДЖЕР}", s["manager_name"]))


def safe_filename(name: str) -> str:
    return re.sub(r'[\\/:*?"<>|]+', "", name).strip()[:150]


def _col_px(col: str) -> int:
    """Ширина колонки в пикселях (как её считает Excel для шрифта Calibri 11)."""
    return int(COLS[col] * 7 + 5)


TABLE_PX = sum(_col_px(c) for c in "BCDEFG")


class _Sheet:
    """Лист КП и инструменты для блоков: текущая строка, шрифты, ячейки, картинки."""

    def __init__(self, ws, s: dict, inv: Invoice, meta: KPMeta, images: dict, blocks_dir: Path | None):
        self.ws, self.s, self.inv, self.meta, self.images = ws, s, inv, meta, images
        self.blocks_dir = blocks_dir
        self.accent = _hex(s["accent_color"])
        self.light = _tint(self.accent, 0.92)
        self.r = 1

    # --- оформление ячеек
    def font(self, size=10, bold=False, color=DARK, italic=False):
        return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)

    def put(self, row, col, value, merge_to=None, f=None, align="left", valign="center", wrap=False,
            fill=None, fmt=None):
        cell = self.ws[f"{col}{row}"]
        cell.value = value
        cell.font = f or self.font()
        cell.alignment = Alignment(horizontal=align, vertical=valign, wrap_text=wrap)
        if fill:
            cell.fill = PatternFill("solid", fgColor=fill)
        if fmt:
            cell.number_format = fmt
        if merge_to:
            self.ws.merge_cells(f"{col}{row}:{merge_to}{row}")
        return cell

    def fill_row(self, row, color, c1=FIRST, c2=LAST):
        for c in range(ord(c1), ord(c2) + 1):
            self.ws[f"{chr(c)}{row}"].fill = PatternFill("solid", fgColor=color)

    def height(self, row, h):
        self.ws.row_dimensions[row].height = h

    def paragraph(self, row, text, f=None, h_per_line=14.5, width=TABLE_WIDTH, align="left", c1=FIRST, c2=LAST):
        self.put(row, c1, text, c2, f=f, wrap=True, valign="top", align=align)
        self.height(row, _lines(text, width * 1.05) * h_per_line + 3)

    # --- картинки
    def place_image(self, path, row: int, max_w, max_h, dy=0, x=0):
        """Вписать картинку в рамку max_w x max_h px; x — отступ в px от левого края колонки B,
        dy — вниз от строки row."""
        img = XLImage(str(path))
        k = min(max_w / img.width, max_h / img.height)
        img.width, img.height = img.width * k, img.height * k
        # Excel не даёт смещению выйти за размер ячейки — переносим его на нужные строку и колонку
        r0 = row - 1
        while True:
            row_px = (self.ws.row_dimensions[r0 + 1].height or 15) * 96 / 72
            if dy < row_px:
                break
            dy -= row_px
            r0 += 1
        cols = "BCDEFGH"
        ci = 0
        while ci < len(cols) - 1 and x >= _col_px(cols[ci]):
            x -= _col_px(cols[ci])
            ci += 1
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(col=column_index_from_string(cols[ci]) - 1, colOff=pixels_to_EMU(x),
                               row=r0, rowOff=pixels_to_EMU(dy)),
            ext=XDRPositiveSize2D(pixels_to_EMU(img.width), pixels_to_EMU(img.height)),
        )
        self.ws.add_image(img)
        return img

    def rich_lines(self, lines) -> list[tuple[list, int]]:
        """Оформленные строки -> [(кусочки (текст, жирный, курсив, цвет, размер), наибольший размер), ...]."""
        colors = {"text": DARK, "muted": GREY, "accent": self.accent}
        out = []
        for line in lines:
            runs, size = [], 10
            for r in line:
                text = fill_placeholders(r["t"], self.s, self.meta)
                if not text:
                    continue
                size = max(size, r["s"])
                runs.append((text, r["b"], r["i"], colors[r["c"]], r["s"]))
            out.append((runs, size))
        return out

    def col_x(self, col: str) -> float:
        """Отступ колонки от левого края колонки B, px."""
        return sum(_col_px(c) for c in "BCDEFG"[: "BCDEFG".index(col)])

    def line_cell(self, row: int, c1: str, c2: str, runs: list, align: str = "left", height: float | None = None):
        """Одна строка текста с оформлением в объединённых ячейках c1:c2.
        Метка {WhatsApp} заменяется местом под значок, а сам значок кладётся поверх ячейки."""
        from core.icons import whatsapp_icon
        blocks, icons, x = [], [], 0.0
        size_max = max([r[4] for r in runs] or [10])
        for text, b, i, color, size in runs:
            parts = WA_TOKEN.split(text)
            for n, part in enumerate(parts):
                if n:  # перед этой частью стояла метка значка
                    icon_px = round(size * 96 / 72 * 1.1)
                    # место под значок — неразрывные пробелы: Excel не схлопывает их и не спотыкается
                    # о кусочек текста из одних обычных пробелов
                    space = _text_px("\u00a0", size, b, i) or 3
                    spaces = "\u00a0" * (int((icon_px + 4) / space) + 1)
                    icons.append((x + 2, icon_px))
                    blocks.append(TextBlock(InlineFont(rFont=FONT, sz=size, b=b, i=i, color=color), spaces))
                    x += _text_px(spaces, size, b, i)
                if part:
                    blocks.append(TextBlock(InlineFont(rFont=FONT, sz=size, b=b, i=i, color=color), part))
                    x += _text_px(part, size, b, i)
        if blocks:
            self.put(row, c1, CellRichText(*blocks), c2, align=align)
        h_pt = height or max(16, size_max * 1.55)
        self.height(row, h_pt)
        if icons:
            area_x = self.col_x(c1)
            area_w = self.col_x(c2) + _col_px(c2) - area_x
            pad = 3
            start = {"left": area_x + pad, "right": area_x + area_w - pad - x,
                     "center": area_x + (area_w - x) / 2}[align]
            row_px = h_pt * 96 / 72
            for ix, icon_px in icons:
                self.place_image(whatsapp_icon(), row, icon_px, icon_px, dy=max(0, (row_px - icon_px) / 2),
                                 x=start + ix)

    def block_image(self, name: str) -> Path | None:
        if not name or not self.blocks_dir:
            return None
        p = self.blocks_dir / name
        return p if p.exists() else None

    @staticmethod
    def image_size(path, max_w, max_h) -> tuple[float, float]:
        from PIL import Image
        with Image.open(path) as im:
            w, h = im.size
        k = min(max_w / w, max_h / h)
        return w * k, h * k

    def x_for(self, align: str, w: float, area_px: float = TABLE_PX, area_x: float = 0) -> float:
        if align == "center":
            return area_x + max(0, (area_px - w) / 2)
        if align == "right":
            return area_x + max(0, area_px - w)
        return area_x


# ---------------------------------------------------------------- стандартные блоки
def _b_header(sh: _Sheet, p: dict):
    s, img = sh.s, sh.images
    if sh.r == 1:
        sh.height(1, 8)
        sh.r = 2
    r = sh.r
    contacts = [
        (s["company_full_name"] or s["company_name"], sh.font(11, True)),
        (s["address"], sh.font(9, color=GREY)),
        ("   ·   ".join(x for x in (s["phone"] and f"Тел.: {s['phone']}", s["email"]) if x),
         sh.font(9, color=GREY)),
        (s["website"], sh.font(9, True, sh.accent)),
    ]
    contacts = [c for c in contacts if c[0]]
    logos_right = p["logo_side"] == "right"
    for i, (text, f) in enumerate(contacts):
        c1, c2, al = (FIRST, "E", "left") if logos_right else ("C", LAST, "right")
        if WA_TOKEN.search(text):
            sh.line_cell(r + i, c1, c2, [(text, bool(f.b), bool(f.i), f.color.rgb[-6:], f.sz)], al, height=16)
            continue
        sh.put(r + i, c1, text, c2, f=f, align=al, wrap=True)
        sh.height(r + i, 16 if _lines(text, TABLE_WIDTH - 20) == 1 else 28)
    rows = max(len(contacts), 4)
    # основной логотип сверху, второй (бренд) под ним
    logo_h = 0
    if img.get("logo"):
        w, _ = sh.image_size(img["logo"], 240, 46)
        logo_h = sh.place_image(img["logo"], r, 240, 46, dy=2, x=sh.x_for("right" if logos_right else "left", w)).height
    if img.get("logo2"):
        w, _ = sh.image_size(img["logo2"], 150, 22)
        sh.place_image(img["logo2"], r, 150, 22, dy=logo_h + 12 if logo_h else 2,
                       x=sh.x_for("right" if logos_right else "left", w))
    sh.r = r + rows
    sh.height(sh.r, 6)
    sh.r += 1
    if p["line"]:
        sh.fill_row(sh.r, sh.accent)
        sh.height(sh.r, 3)
        sh.r += 1
    sh.r += 1


def _b_title(sh: _Sheet, p: dict):
    sh.put(sh.r, FIRST, p["text"], LAST, f=sh.font(p["size"], True, sh.accent), align=p["align"])
    sh.height(sh.r, p["size"] * 1.5)
    sh.r += 1
    if p["show_number"]:
        m = sh.meta
        sub = f"№ {m.number} от {m.date:%d.%m.%Y}"
        if sh.s.get("validity_days"):
            sub += f"   ·   действительно до {m.date + timedelta(days=int(sh.s['validity_days'])):%d.%m.%Y}"
        sh.put(sh.r, FIRST, sub, LAST, f=sh.font(10, color=GREY), align=p["align"])
        sh.height(sh.r, 16)
        sh.r += 1
    sh.r += 1


def _b_parties(sh: _Sheet, p: dict):
    def labeled(row, label, value, details=""):
        rich = CellRichText(
            TextBlock(InlineFont(rFont=FONT, sz=10, color=GREY), f"{label}  "),
            TextBlock(InlineFont(rFont=FONT, sz=10, b=True, color=DARK), value),
        )
        if details:
            rich.append(TextBlock(InlineFont(rFont=FONT, sz=9, color=GREY), f"\n{details}"))
        sh.put(row, FIRST, rich, LAST, wrap=True, valign="top", fill=sh.light)
        sh.fill_row(row, sh.light)
        sh.height(row, _lines(label + value, TABLE_WIDTH) * 15
                  + (_lines(details, TABLE_WIDTH * 1.1) * 13 if details else 0) + 6)

    m, s = sh.meta, sh.s
    start = sh.r
    if p["show_to"] and (m.buyer or m.buyer_details):
        labeled(sh.r, "Кому:", m.buyer or "—", m.buyer_details)
        sh.r += 1
    if p["show_from"] and s["company_name"]:
        labeled(sh.r, "От:", s["company_name"], fill_placeholders(s.get("from_details", ""), s, sh.meta).strip())
        sh.r += 1
    if sh.r > start:
        sh.r += 1


def _b_greeting(sh: _Sheet, p: dict):
    if not richtext.is_empty(sh.s.get("greeting_rich")):
        # приветствие с оформлением — одна ячейка, строки через перенос
        lines = sh.rich_lines(richtext.normalize(sh.s["greeting_rich"]))
        blocks, height = [], 0
        for n, (runs, size) in enumerate(lines):
            if n:
                blocks.append(TextBlock(InlineFont(rFont=FONT, sz=size), NL_CHAR))
            for text, b, i, color, sz in runs:
                blocks.append(TextBlock(InlineFont(rFont=FONT, sz=sz, b=b, i=i, color=color),
                                        WA_TOKEN.sub("WhatsApp", text)))
            line_text = "".join(r[0] for r in runs)
            height += _lines(line_text, TABLE_WIDTH * 10.5 / size) * size * 1.4
        if blocks:
            sh.put(sh.r, FIRST, CellRichText(*blocks), LAST, wrap=True, valign="top")
            sh.height(sh.r, height + 3)
            sh.r += 2
        return
    greeting = fill_placeholders(sh.s["greeting"], sh.s, sh.meta).strip()
    if greeting:
        size = 10.5 if p["size"] == 10 else p["size"]  # 10 — исторически 10.5 pt
        sh.paragraph(sh.r, greeting, sh.font(size), h_per_line=14.5 * size / 10.5, width=TABLE_WIDTH * 10.5 / size)
        sh.r += 2


def _b_table(sh: _Sheet, p: dict):
    thin = Side(style="thin", color=LINE)
    headers = [("B", "№", "center"), ("C", "Наименование", "left"), ("D", "Кол-во", "center"),
               ("E", "Ед.", "center"), ("F", "Цена", "right"), ("G", "Сумма", "right")]
    for col, title, al in headers:
        sh.put(sh.r, col, title, f=sh.font(10, True, "FFFFFF"), align=al, fill=sh.accent)
    sh.height(sh.r, 22)
    sh.r += 1
    for i, it in enumerate(sh.inv.items, 1):
        zebra = "FFFFFF" if i % 2 else "F9FAFB"
        vals = [("B", i, "center", None), ("C", it.name, "left", None),
                ("D", _fmt_qty(it.qty), "center", None), ("E", it.unit, "center", None),
                ("F", it.price, "right", "#,##0.00"), ("G", it.total, "right", "#,##0.00")]
        for col, v, al, fmt in vals:
            c = sh.put(sh.r, col, v, align=al, wrap=(col == "C"), fill=zebra, fmt=fmt,
                       f=sh.font(10, bold=(col == "G")))
            c.border = Border(bottom=thin)
        sh.height(sh.r, max(20, _lines(it.name, COLS["C"] * 1.12) * 14 + 6))
        sh.r += 1
    sh.r += 1


def _b_totals(sh: _Sheet, p: dict):
    def total_row(label, value, big=False):
        sh.put(sh.r, "D", label, "F", f=sh.font(12 if big else 10, big, DARK if big else GREY), align="right")
        sh.put(sh.r, "G", value, f=sh.font(12 if big else 10, True, sh.accent if big else DARK), align="right",
               fmt="#,##0.00")
        if big:
            sh.fill_row(sh.r, sh.light, "D", "G")
        sh.height(sh.r, 24 if big else 17)
        sh.r += 1

    inv = sh.inv
    if sh.s.get("show_vat", True) and inv.vat is not None:
        total_row("Итого:", inv.total)
        total_row(inv.vat_label.rstrip(":") + ":", inv.vat)
    total_row("Сумма предложения, KZT:", inv.total, big=True)
    sh.r += 1


def _b_amount_words(sh: _Sheet, p: dict):
    inv = sh.inv
    sh.paragraph(sh.r, f"Всего наименований {len(inv.items)}, на сумму {_fmt_money(inv.total)} KZT\n"
                       f"{amount_in_words(inv.total)}", sh.font(10, italic=True, color=GREY))
    sh.r += 2


def _b_conditions(sh: _Sheet, p: dict):
    conditions = [c.strip() for c in sh.s["conditions"].splitlines() if c.strip()]
    if not conditions:
        return
    if p["heading"]:
        sh.put(sh.r, FIRST, p["heading"], LAST, f=sh.font(12, True, sh.accent))
        sh.height(sh.r, 20)
        for c in range(ord(FIRST), ord(LAST) + 1):
            sh.ws[f"{chr(c)}{sh.r}"].border = Border(bottom=Side(style="thin", color=sh.accent))
        sh.r += 1
    for line in conditions:
        sh.put(sh.r, FIRST, "•", f=sh.font(10, True, sh.accent), align="center", valign="top")
        sh.put(sh.r, "C", line, LAST, wrap=True, valign="top")
        sh.height(sh.r, _lines(line, (TABLE_WIDTH - COLS["B"]) * 1.25) * 14.5 + 2)
        sh.r += 1
    sh.r += 1


def signature_lines(s: dict) -> list:
    """Строки подписи с оформлением: заданные в «Текстах» или собранные из контактов."""
    if not richtext.is_empty(s.get("signature_rich")):
        return richtext.normalize(s["signature_rich"])
    text = s.get("signature_text", "").strip() or auto_signature(s)
    # прежнее оформление: первая строка серая, вторая жирная, остальные серые
    return richtext.from_plain(text, lambda i: dict(b=True, s=11) if i == 1 else dict(c="muted", s=10))


def auto_signature(s: dict) -> str:
    """Подпись по умолчанию — из контактов компании (если текст подписи не задан в «Текстах»)."""
    lines = ["С уважением,"]
    if s["manager_name"]:
        lines.append(s["manager_name"])
        if s["manager_position"]:
            lines.append(f"{s['manager_position']}, {s['company_name']}")
    elif s["company_name"]:
        lines.append(s["company_name"])
    if s["manager_phone"]:
        lines.append(f"Тел.: {s['manager_phone']}")
    if s["email"]:
        lines.append(s["email"])
    return "\n".join(lines)


def _b_signature(sh: _Sheet, p: dict):
    s = sh.s
    stamp_left = p["stamp_side"] == "left"
    c1, c2 = ("D", LAST) if stamp_left else (FIRST, "C")
    start = sh.r
    for runs, size in sh.rich_lines(signature_lines(s)):
        sh.line_cell(sh.r, c1, c2, runs)
        sh.r += 1
    if s["show_stamp"] and sh.images.get("stamp"):
        w, _ = sh.image_size(sh.images["stamp"], 260, 115)
        # печать прижимаем к краю страницы: слева — к левому, справа — к правому
        sh.place_image(sh.images["stamp"], max(start - 1, 1), 260, 115, x=sh.x_for("left" if stamp_left else "right", w))
        sh.r = max(sh.r, start + 6)
    sh.r += 1


# ---------------------------------------------------------------- свои блоки
_TEXT_COLORS = {"text": DARK, "muted": GREY}


def _b_text(sh: _Sheet, p: dict):
    text = p["text"].strip()
    if not text:
        return
    color = sh.accent if p["color"] == "accent" else _TEXT_COLORS[p["color"]]
    f = sh.font(p["size"], p["bold"], color, p["italic"])
    sh.paragraph(sh.r, text, f, h_per_line=p["size"] * 1.45, width=TABLE_WIDTH * 10 / p["size"], align=p["align"])
    if p["background"]:
        sh.fill_row(sh.r, sh.light)
        sh.height(sh.r, (sh.ws.row_dimensions[sh.r].height or 15) + 8)
        sh.ws[f"{FIRST}{sh.r}"].alignment = Alignment(horizontal=p["align"], vertical="center", wrap_text=True,
                                                     indent=1 if p["align"] == "left" else 0)
    sh.r += 2


def _b_image(sh: _Sheet, p: dict):
    path = sh.block_image(p["image"])
    if not path:
        return
    w, h = sh.image_size(path, min(p["width"], TABLE_PX), 2000)
    sh.height(sh.r, h * 72 / 96 + 6)
    sh.place_image(path, sh.r, w, h, dy=3, x=sh.x_for(p["align"], w))
    sh.r += 2


def _b_columns(sh: _Sheet, p: dict):
    left_px = _col_px("B") + _col_px("C")
    right_px = TABLE_PX - left_px
    heights = [16]
    for side, c1, c2, area_x, area_w, chars in (("left", FIRST, "C", 0, left_px, COLS["B"] + COLS["C"]),
                                                ("right", "D", LAST, left_px + 10, right_px - 10,
                                                 sum(COLS[c] for c in "DEFG"))):
        path = sh.block_image(p[f"{side}_image"])
        if path:
            w, h = sh.image_size(path, min(p["image_width"], area_w), 600)
            sh.place_image(path, sh.r, w, h, dy=3, x=area_x)
            heights.append(h * 72 / 96 + 6)
        elif p[f"{side}_text"].strip():
            text = p[f"{side}_text"].strip()
            sh.put(sh.r, c1, text, c2, f=sh.font(p["size"]), wrap=True, valign="top")
            heights.append(_lines(text, chars * 10 / p["size"] * 1.05) * p["size"] * 1.45 + 3)
    sh.height(sh.r, max(heights))
    sh.r += 2


def _b_divider(sh: _Sheet, p: dict):
    sh.fill_row(sh.r, sh.accent if p["color"] == "accent" else "D1D5DB")
    sh.height(sh.r, 3 if p["thickness"] == "thick" else 1.5)
    sh.r += 2


def _b_spacer(sh: _Sheet, p: dict):
    sh.height(sh.r, p["height"])
    sh.r += 1


RENDERERS = {
    "header": _b_header, "title": _b_title, "parties": _b_parties, "greeting": _b_greeting,
    "table": _b_table, "totals": _b_totals, "amount_words": _b_amount_words, "conditions": _b_conditions,
    "signature": _b_signature, "text": _b_text, "image": _b_image, "columns": _b_columns,
    "divider": _b_divider, "spacer": _b_spacer,
}


def build_kp(inv: Invoice, meta: KPMeta, s: dict, images: dict[str, Path | None],
             out_path: Path, blocks_dir: Path | None = None) -> Path:
    """images: {'logo': ..., 'logo2': ..., 'stamp': ...}, None — картинки нет.
    blocks_dir — папка с картинками своих блоков (data/blocks)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "КП"
    ws.sheet_view.showGridLines = False
    for col, w in COLS.items():
        ws.column_dimensions[col].width = w

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.4, bottom=0.5)
    ws.print_options.horizontalCentered = True
    # Подвал печатается внизу каждой страницы и не сдвигает содержимое
    footer_text = " ".join(s["footer"].split()).replace("&", "&&")[:200]
    if footer_text:
        ws.oddFooter.left.text = footer_text
        ws.oddFooter.left.size = 7
        ws.oddFooter.left.color = GREY
    ws.oddFooter.right.text = "стр. &P из &N"
    ws.oddFooter.right.size = 7
    ws.oddFooter.right.color = GREY

    sh = _Sheet(ws, s, inv, meta, images, blocks_dir)
    for block in normalize(s.get("layout"), s):
        if block["visible"]:
            RENDERERS[block["type"]](sh, block["props"])

    ws.print_area = f"A1:H{max(sh.r, 2)}"
    _apply_legal_tenge(ws)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def _apply_legal_tenge(ws):
    """По закону слово пишется «теңге» и в русском тексте — правим все ячейки и подвал."""
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str):
                cell.value = legal_tenge(v)
            elif isinstance(v, CellRichText):
                for part in v:
                    if isinstance(part, TextBlock):
                        part.text = legal_tenge(part.text)
    for hf in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
        if hf.text:
            hf.text = legal_tenge(hf.text)


def export_pdf(xlsx: Path) -> Path:
    """PDF через установленный Excel. Бросает исключение, если Excel недоступен."""
    import pythoncom
    import win32com.client
    pdf = xlsx.with_suffix(".pdf")
    pythoncom.CoInitialize()
    excel = win32com.client.DispatchEx("Excel.Application")
    excel.Visible = False
    excel.DisplayAlerts = False
    try:
        wb = excel.Workbooks.Open(str(xlsx.resolve()), ReadOnly=True)
        wb.ExportAsFixedFormat(0, str(pdf.resolve()))
        wb.Close(False)
    finally:
        excel.Quit()
        pythoncom.CoUninitialize()
    return pdf


def demo_invoice(settings: dict | None = None) -> tuple[Invoice, KPMeta]:
    """Данные для предпросмотра в редакторе шаблона.

    Только вымышленные обезличенные позиции — без брендов и без данных из реальных документов.
    """
    inv = Invoice(
        number="0000", date=date.today(), buyer="ТОО «Пример покупателя»",
        buyer_details="БИН / ИИН 000000000000, г. Город, ул. Примерная, 1",
        items=[
            Item("Пример товара № 1", 1, "шт", 10000, 10000),
            Item("Пример товара № 2", 2, "шт", 2500, 5000),
            Item("Пример товара № 3 с длинным названием, чтобы проверить перенос строк "
                 "в таблице коммерческого предложения", 3, "шт", 1500, 4500),
        ],
        total=19500, vat=round(19500 * 16 / 116, 2), vat_label="В т.ч. НДС (16%):",
    )
    # предпросмотр показывает КП так же, как настоящий: с учётом переключателей «Из счёта»
    s = settings or {}
    return inv, KPMeta(number="0000", date=date.today(),
                       buyer=inv.buyer if s.get("fill_buyer", True) else "",
                       buyer_details=inv.buyer_details if s.get("fill_buyer_details", True) else "")
